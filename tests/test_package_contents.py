"""Shared B2B package details and unit-safe historical volume upgrades."""
from decimal import Decimal
import io
import json

import pytest
from openpyxl import load_workbook

from app import orm as models
from app.models import db
from label_helpers import label_form
from test_label_review import panel, create, pages, snapshot, migrated_admin, populated, legacy_engine, disposable_engine
from test_complete_backup import unpack


CONTENTS = {'kind': 'box', 'default_pieces': 10}
OLD_CONTENTS = {**CONTENTS,
               'boxes': {'2': {'pieces': 8, 'note': 'Fragile & upright <glass>', 'weight': '3.250'}}}


def saved(client, **values):
    data = {'consignment_number': 'BOX-CONTENTS', 'pieces': 3, 'package_contents': CONTENTS,
            'measurement_mode': 'volume', 'volume': '25.250', **values}
    response = client.post('/admin/consignments/save', json={'rows': [data]})
    assert response.status_code == 200, response.json
    return client.get('/admin/consignments/list?search=BOX-CONTENTS').json['rows'][0]


def test_shared_counts_persist_and_print_on_every_page(panel):
    app, client = panel
    row = saved(client)
    assert row['package_contents'] == CONTENTS
    assert row['pieces'] == 3 and row['volume_unit'] == 'ft3'
    before = snapshot(app, include_label_runs=False)
    review = client.post('/admin/labels/review', json={'consignment_ids': [row['id']]}).json
    assert review['shipments'][0]['package_contents'] == CONTENTS
    output = pages(client.post('/admin/labels/generate', data=label_form([row['id']])))
    assert len(output) == 3
    for index, page in enumerate(output, 1):
        text = page.extract_text()
        assert 'PIECES\n' in text and 'BOX\n' in text and 'BOXES\n' not in text
        assert f'{index} of 3' in text and '25.250 ft³' in text and 'm³' not in text
        assert '\n10\n' in text
    assert snapshot(app, include_label_runs=False) == before
    # Older callers that omit contents must not erase them.
    result = client.post('/admin/consignments/save', json={'rows': [{'id': row['id'], 'consignment_number': row['consignment_number'], 'status': 'Delivered'}]})
    assert result.status_code == 200
    assert client.get('/admin/consignments/list?search=BOX-CONTENTS').json['rows'][0]['package_contents'] == CONTENTS


def test_saved_exceptions_cannot_affect_reads_exports_or_new_labels(panel):
    app, client = panel
    row = saved(client, measurement_mode='weight', weight='12.345')
    # Simulate an existing database containing data saved before the removal.
    with app.app_context():
        stored = db.session.get(models.Consignment, models.record_id(row['id']))
        stored.package_contents = OLD_CONTENTS
        db.session.commit()
    before = snapshot(app, include_label_runs=False)
    listed = client.get('/admin/consignments/list?search=BOX-CONTENTS').json['rows'][0]
    assert listed['package_contents'] == CONTENTS
    reviewed = client.post('/admin/labels/review', json={'consignment_ids': [row['id']]}).json['shipments'][0]
    assert reviewed['package_contents'] == CONTENTS and 'boxes' not in reviewed
    sheet = list(load_workbook(io.BytesIO(client.get('/admin/consignments/export.xlsx').data)).active.values)
    exported = next(dict(zip(sheet[0], values)) for values in sheet[1:] if values[0] == 'BOX-CONTENTS')
    assert json.loads(exported['package_contents']) == CONTENTS
    data = label_form([row['id']])
    details = json.loads(data['label_details'])
    details[0]['box_invoice_amounts'] = ['6000.00', '50.25', '7000']
    data['label_details'] = json.dumps(details)
    output = pages(client.post('/admin/labels/generate', data=data))
    assert len(output) == 3
    for page in output:
        text = page.extract_text()
        assert '12.345 kg' in text and 'Rs 6,000.00' in text and '\n10\n' in text
        for retired in ('3.250 kg', 'Rs 50.25', 'Rs 7,000.00', 'Fragile', '\n8\n', 'BOX NOTE'):
            assert retired not in text
    assert snapshot(app, include_label_runs=False) == before, 'Reads and printing must not rewrite saved data'
    # Saving the record removes retired overrides, including for older clients.
    assert client.post('/admin/consignments/save', json={'rows': [{**listed, 'package_contents': OLD_CONTENTS}]}).status_code == 200
    with app.app_context():
        assert db.session.get(models.Consignment, models.record_id(row['id'])).package_contents == CONTENTS


@pytest.mark.parametrize('contents', [
    {'kind': 'bag'}, {'default_pieces': -1}, {'default_pieces': True}, {'default_pieces': '1.5'},
    {'default_pieces': 0}, {'default_pieces': 1000001}, {'unknown': 2}, [],
])
def test_bad_contents_rejected_without_partial_save(panel, contents):
    app, client = panel
    before = snapshot(app)
    response = client.post('/admin/consignments/save', json={'rows': [{'consignment_number': 'INVALID-CONTENTS', 'pieces': 3, 'package_contents': contents}]})
    assert response.status_code == 400
    assert snapshot(app) == before


def test_reducing_boxes_discards_old_overrides_without_blocking_the_save(panel):
    _, client = panel
    row = saved(client)
    response = client.post('/admin/consignments/save', json={'rows': [{**row, 'pieces': 1, 'package_contents': OLD_CONTENTS}]})
    assert response.status_code == 200
    stored = client.get('/admin/consignments/list?search=BOX-CONTENTS').json['rows'][0]
    assert stored['pieces'] == 1 and stored['package_contents'] == CONTENTS


def test_unknown_piece_counts_and_pallets_are_not_invented(panel):
    _, client = panel
    row = saved(client, package_contents={'kind': 'pallet', 'default_pieces': None, 'boxes': {}})
    for page in pages(client.post('/admin/labels/generate', data=label_form([row['id']]))):
        text = page.extract_text()
        assert 'PALLET\n' in text and 'PIECES\n' in text and '—' in text


def test_old_cbm_is_displayed_correctly_and_upgrade_converts_exactly_once(panel):
    app, client = panel
    row = saved(client)
    with app.app_context():
        stored = db.session.get(models.Consignment, models.record_id(row['id']))
        stored.volume = Decimal('1.000')
        stored.volume_unit = None
        db.session.commit()
    before = snapshot(app)
    converted = client.get('/admin/consignments/list?search=BOX-CONTENTS').json['rows'][0]
    assert converted['volume'] == '35.315' and converted['volume_unit'] == 'ft3'
    assert snapshot(app) == before, 'Read-only views must not mutate the database'
    from app.shipment_packages import upgrade_volume_units
    for attempt in range(2):
        with app.app_context():
            assert upgrade_volume_units() == (1 if attempt == 0 else 0)
            stored = db.session.get(models.Consignment, models.record_id(row['id']))
            assert stored.volume == Decimal('35.315') and stored.volume_unit == 'ft3'
    result = client.post('/admin/consignments/save', json={'rows': [converted]})
    assert result.status_code == 200
    assert client.get('/admin/consignments/list?search=BOX-CONTENTS').json['rows'][0]['volume'] == '35.315'


def test_box_contents_and_ft3_survive_excel_and_complete_backup(panel):
    app, client = panel
    saved(client)
    sheet = load_workbook(io.BytesIO(client.get('/admin/consignments/export.xlsx').data)).active
    rows = list(sheet.values)
    headers = rows[0]
    record = next(dict(zip(headers, values)) for values in rows[1:] if values[0] == 'BOX-CONTENTS')
    assert record['boxes_pallets'] == 3 and record['volume_unit'] == 'ft3' and record['volume'] == 25.25
    assert json.loads(record['package_contents']) == CONTENTS
    from openpyxl import Workbook
    book = Workbook(); book.active.append(headers)
    record['consignment_number'] = 'BOX-CONTENTS-IMPORTED'
    record['package_contents'] = json.dumps(OLD_CONTENTS)
    book.active.append([record.get(key) for key in headers])
    content = io.BytesIO(); book.save(content)
    response = client.post('/admin/consignments/import', data={'file': (io.BytesIO(content.getvalue()), 'contents.xlsx')}, follow_redirects=True)
    assert response.status_code == 200 and b'Added: 1' in response.data
    imported = client.get('/admin/consignments/list?search=BOX-CONTENTS-IMPORTED').json['rows'][0]
    assert imported['package_contents'] == CONTENTS and imported['volume'] == '25.250'
    archive, payload, _ = unpack(client.get('/admin/generate-backup'))
    workbook = load_workbook(io.BytesIO(archive.read('admin-data.xlsx')))
    values = list(workbook['consignments'].values)
    records = [dict(zip(values[0], row)) for row in values[1:]]
    backed_up = next(row for row in records if row.get('consignment_number', row.get('identifier_value')) == 'BOX-CONTENTS')
    assert json.loads(backed_up['package_contents']) == CONTENTS and backed_up['volume_unit'] == 'ft3'


def test_volume_upgrade_rolls_back_all_rows_if_one_value_overflows(panel):
    app, client = panel
    row = saved(client)
    with app.app_context():
        stored = db.session.get(models.Consignment, models.record_id(row['id']))
        stored.volume = Decimal('999999999.999')
        stored.volume_unit = None
        db.session.commit()
    before = snapshot(app)
    from app.shipment_packages import upgrade_volume_units
    with app.app_context(), pytest.raises(ValueError, match='outside the supported range'):
        upgrade_volume_units()
    assert snapshot(app) == before


def test_complete_database_migration_keeps_contents_and_volume_unit(panel, isolated_storage, tmp_path):
    app, client = panel
    saved(client)
    from app.er_adapter import prepare_backend, RecordExtras
    from database.models import Consignment
    from sqlalchemy import select
    from sqlalchemy.orm import Session
    from migrations.shipment_v1 import migrate
    from migrations.storage import Storage
    with app.app_context():
        source = db.engine
        with disposable_engine(source.dialect.name, 'copied-boxes') as target:
            migrate(target, Storage(isolated_storage), source_engine=source, fresh_target=True)
            prepare_backend(target)
            with Session(target) as session:
                record = session.scalar(select(Consignment).where(Consignment.identifier_value == 'BOX-CONTENTS'))
                extras = session.get(RecordExtras, ('consignments', record.id)).data
                assert extras['volume_unit'] == 'ft3' and extras['package_contents'] == CONTENTS
                assert record.pieces == 3 and record.volume == Decimal('25.250')
