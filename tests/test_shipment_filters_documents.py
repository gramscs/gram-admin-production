"""Combined shipment filters and the document library on both database models."""
import io
import json
import re
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from test_bulk_documents import image, pdf, upload
from test_er_dashboard import migrated_admin, populated, legacy_engine, disposable_engine


def dataset(client):
    a = client.post('/admin/companies', json={'name': 'Filter Client A'}).json['company']
    b = client.post('/admin/companies', json={'name': 'Filter Client B'}).json['company']
    rows = [
        ('FILTER-A-LRN', 'LRN', a['id']), ('FILTER-A1', 'AWB', a['id']),
        ('FILTER-A2', 'AWB', a['id']), ('FILTER-A-ORDER', 'Order ID', a['id']),
        ('FILTER-B', 'AWB', b['id']), ('FILTER-NONE', 'AWB', None),
    ]
    result = client.post('/admin/consignments/save', json={'rows': [
        {'consignment_number': number, 'identifier_type': kind, 'company_id': company,
         'status': 'In Transit', 'pieces': 1} for number, kind, company in rows]})
    assert result.status_code == 200, result.json
    assert client.post('/admin/companies/' + str(b['id']) + '/archive', json={}).status_code == 200
    return a, b


def combined_filters(client):
    a, b = dataset(client)
    filters = {'company': str(a['id']), 'identifier_type': 'AWB', 'search': 'FILTER-', 'per_page': 1}
    first = client.get('/admin/consignments/list', query_string=filters)
    assert first.status_code == 200, first.json
    assert first.json['total'] == 2 and first.json['pages'] == 2 and first.json['has_next']
    second = client.get('/admin/consignments/list', query_string={**filters, 'page': 2})
    assert second.json['has_prev'] and not second.json['has_next']
    assert {first.json['rows'][0]['consignment_number'], second.json['rows'][0]['consignment_number']} == {'FILTER-A1', 'FILTER-A2'}
    for values, expected in [
        ({'company': str(a['id'])}, 4), ({'identifier_type': 'AWB'}, 4),
        ({'company': 'unassigned', 'identifier_type': 'AWB'}, 1),
        ({'company': str(b['id']), 'identifier_type': 'AWB'}, 1),
        ({'company': str(a['id']), 'identifier_type': 'Order ID'}, 1),
        ({'company': str(a['id']), 'identifier_type': 'AWB', 'search': 'FILTER-A1'}, 1),
    ]:
        result = client.get('/admin/consignments/list', query_string={'search': 'FILTER-', **values})
        assert result.status_code == 200 and result.json['total'] == expected, result.json


def filtered_exports(client):
    a, b = dataset(client)
    filters = {'company': str(a['id']), 'identifier_type': 'AWB', 'search': 'FILTER-'}
    excel = client.get('/admin/consignments/export.xlsx', query_string=filters)
    assert excel.status_code == 200
    rows = list(load_workbook(io.BytesIO(excel.data)).active.values)
    assert {row[0] for row in rows[1:]} == {'FILTER-A1', 'FILTER-A2'}
    exported = client.get('/admin/consignments/export.pdf', query_string=filters)
    assert exported.status_code == 200
    text = '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(exported.data)).pages)
    assert 'FILTER-A1' in text and 'FILTER-A2' in text
    assert 'FILTER-A-LRN' not in text and 'FILTER-A-ORDER' not in text and 'FILTER-B' not in text and 'FILTER-NONE' not in text
    empty = client.get('/admin/consignments/export.xlsx', query_string={**filters, 'search': 'NO-MATCH'})
    assert load_workbook(io.BytesIO(empty.data)).active.max_row == 1


def filtered_document_zip(client):
    a, b = dataset(client)
    files = [('FILTER-A1_pod.png', image()), ('FILTER-A1_invoice.pdf', pdf()),
             ('FILTER-A2_pod.png', image()), ('FILTER-B_pod.png', image()),
             ('FILTER-A-LRN_invoice.pdf', pdf())]
    assert upload(client, files, apply=True).status_code == 200
    result = client.post('/admin/consignments/documents/bulk-download', json={
        'company': str(a['id']), 'identifier_type': 'AWB', 'search': 'FILTER-', 'kind': 'both'})
    assert result.status_code == 200
    with ZipFile(io.BytesIO(result.data)) as archive:
        assert set(archive.namelist()) == {'FILTER-A1_pod.png', 'FILTER-A1_invoice.pdf', 'FILTER-A2_pod.png', 'document-manifest.json'}
        manifest = json.loads(archive.read('document-manifest.json'))
        assert manifest['shipment_count'] == 2 and manifest['downloaded'] == 3
        assert all(row['identifier_type'] == 'AWB' and row['company'] == a['name'] for row in manifest['documents'])
    # Explicit selections can span filters; they remain an intentional scope.
    selected = client.get('/admin/consignments/list', query_string={'search': 'FILTER-B'}).json['rows'][0]
    result = client.post('/admin/consignments/documents/bulk-download', json={
        'ids': [selected['id']], 'company': str(a['id']), 'identifier_type': 'LRN', 'kind': 'pod'})
    assert result.status_code == 200
    with ZipFile(io.BytesIO(result.data)) as archive:
        assert 'FILTER-B_pod.png' in archive.namelist()


def library_and_linked_details(client):
    a, b = dataset(client)
    for path in ('/admin/shipments', '/admin/consignments', '/admin/documents'):
        result = client.get(path, query_string={'company': str(b['id']), 'identifier_type': 'AWB'})
        assert result.status_code == 200
        assert b'>Shipments</span>' in result.data and b'>Documents</span>' in result.data
        assert b'id="company-filter"' in result.data and b'id="identifier-type-filter"' in result.data
        assert b'Filter Client B (archived)' in result.data
        assert b'id="bulk-document-files"' in result.data and b'id="bulk-review-btn"' in result.data
    assert upload(client, [('FILTER-A1_pod.png', image()), ('FILTER-A1_invoice.pdf', pdf())], apply=True).status_code == 200
    row = client.get('/admin/consignments/list', query_string={'search': 'FILTER-A1'}).json['rows'][0]
    assert row['company_name'] == a['name'] and row['identifier_type'] == 'AWB'
    assert row['pod_original_name'] == 'FILTER-A1_pod.png' and row['invoice_original_name'] == 'FILTER-A1_invoice.pdf'
    changed = client.post('/admin/consignments/save', json={'rows': [{**row, 'identifier_type': 'Order ID'}]})
    assert changed.status_code == 200, changed.json
    updated = client.get('/admin/consignments/list', query_string={'search': 'FILTER-A1', 'identifier_type': 'Order ID'}).json['rows'][0]
    assert updated['pod_image'] == row['pod_image'] and updated['invoice_file'] == row['invoice_file']
    assert client.get('/admin/consignments/list', query_string={'search': 'FILTER-A1', 'identifier_type': 'AWB'}).json['total'] == 0


def invalid_filters(client):
    for filters in ({'identifier_type': 'ISBN'}, {'company': 'unknown-company'}, {'search': 'x' * 201}):
        for path in ('/admin/consignments/list', '/admin/consignments/export.xlsx', '/admin/consignments/export.pdf', '/admin/documents', '/admin/shipments'):
            assert client.get(path, query_string=filters).status_code == 400, (path, filters)
        assert client.post('/admin/consignments/documents/bulk-download', json=filters).status_code == 400
    for filters in ({'identifier_type': []}, {'company': True}, {'search': 123}):
        assert client.post('/admin/consignments/documents/bulk-download', json=filters).status_code == 400


def saved_view_identifier_filter(client):
    a, b = dataset(client)
    settings = {'name': 'AWB only', 'format': 'csv', 'columns': ['consignment_number', 'identifier_type'],
                'filters': {'period': 'all', 'company': str(a['id']), 'identifier_type': 'AWB'}}
    saved = client.post('/admin/mis/views', json=settings)
    assert saved.status_code == 200 and saved.json['view']['filters']['identifier_type'] == 'AWB'
    generated = client.post('/admin/mis/reports', json={**settings, 'view_id': saved.json['view']['id']})
    assert generated.status_code == 200, generated.json
    content = client.get(generated.json['download_url']).data.decode('utf-8-sig')
    assert 'FILTER-A1' in content and 'FILTER-A2' in content and 'FILTER-A-LRN' not in content and 'FILTER-B' not in content


def dashboard_missing_documents(client):
    a, b = dataset(client)
    assert upload(client, [('FILTER-A1_pod.png', image()), ('FILTER-A2_invoice.pdf', pdf()),
                          ('FILTER-A-LRN_pod.png', image()), ('FILTER-A-LRN_invoice.pdf', pdf())], apply=True).status_code == 200
    def dashboard(filters):
        response = client.get('/admin/dashboard', query_string=filters)
        assert response.status_code == 200
        html = response.get_data(as_text=True)
        match = re.search(r'<script type="application/json" id="missing-document-data">(.*?)</script>', html, re.S)
        assert match and 'follows your dashboard filters' in html
        assert 'data-missing-filter="pod"' in html and 'data-missing-filter="both"' in html
        return json.loads(match.group(1)), html
    rows, html = dashboard({'search': 'FILTER-'})
    assert len(rows) == 5
    assert 'FILTER-A-LRN' not in {row['identifier'] for row in rows}
    assert sum(row['missing_pod'] for row in rows) == 4
    assert sum(row['missing_invoice'] for row in rows) == 4
    assert sum(row['missing_pod'] and row['missing_invoice'] for row in rows) == 3
    assert {row['identifier_type'] for row in rows} == {'AWB', 'Order ID'}
    # Existing delivered-without-POD metric is distinct from all-status follow-up.
    assert all(row['status'] == 'In Transit' for row in rows)
    rows, html = dashboard({'search': 'FILTER-', 'company': a['id'], 'identifier_type': 'AWB'})
    assert {row['identifier'] for row in rows} == {'FILTER-A1', 'FILTER-A2'}
    assert all(row['company_id'] == str(a['id']) and row['company'] == a['name'] for row in rows)
    assert 'Open Documents' in html
    rows, _ = dashboard({'search': 'FILTER-', 'company': 'unassigned'})
    assert len(rows) == 1 and rows[0]['company_id'] == 'unassigned'
    # Date and status scope is the same as dashboard charts, including invalid dates.
    original = client.get('/admin/consignments/list', query_string={'search': 'FILTER-', 'per_page': 100}).json['rows']
    for row in original:
        row['pickup_date'] = '2026-10-01' if row['consignment_number'] == 'FILTER-A1' else '2026-09-01'
        row['status'] = 'Delivered' if row['consignment_number'] == 'FILTER-A1' else 'In Transit'
    assert client.post('/admin/consignments/save', json={'rows': original}).status_code == 200
    rows, _ = dashboard({'search': 'FILTER-', 'period': 'custom', 'start': '2026-10-01', 'end': '2026-10-10', 'status': 'Delivered'})
    assert len(rows) == 1 and rows[0]['identifier'] == 'FILTER-A1' and rows[0]['missing_invoice'] and not rows[0]['missing_pod']
    assert upload(client, [('FILTER-A1_invoice.pdf', pdf())], apply=True).status_code == 200
    rows, html = dashboard({'search': 'FILTER-A1'})
    assert rows == [] and 'All PODs and invoices are uploaded in this view.' in html


def individual_document_management(client):
    a, b = dataset(client)
    row = client.get('/admin/consignments/list', query_string={'search': 'FILTER-A1'}).json['rows'][0]
    assert client.post('/admin/consignments/save', json={'rows': [{
        'consignment_number': 'FILTER-A10', 'identifier_type': 'AWB', 'company_id': a['id'], 'pieces': 3}]}).status_code == 200
    assert client.get('/admin/consignments/list', query_string={'search': 'FILTER-A1'}).json['total'] == 2
    exact = client.get('/admin/consignments/list', query_string={'search': 'FILTER-A1', 'shipment_id': row['id']})
    assert exact.status_code == 200 and exact.json['total'] == 1 and exact.json['rows'][0]['id'] == row['id']
    assert client.get('/admin/consignments/list', query_string={'shipment_id': 'invalid'}).status_code == 400
    html = client.get('/admin/documents').get_data(as_text=True)
    assert 'documentEditorModal' in html and 'document-editor-file' in html and 'document-editor-save' in html
    def current():
        return client.get('/admin/consignments/list', query_string={'shipment_id': row['id']}).json['rows'][0]
    ignored = {'pod_image', 'pod_original_name', 'invoice_file', 'invoice_original_name'}
    original = {key: value for key, value in row.items() if key not in ignored}
    for kind, field in [('pod', 'pod_image'), ('invoice', 'invoice_file')]:
        path = '/admin/consignments/' + str(row['id']) + '/' + kind
        first = client.post(path, data={'expected_reference': '', 'file': (io.BytesIO(pdf()), 'first.pdf')})
        assert first.status_code == 200 and first.json[field]
        old = first.json[field]
        assert client.get(path).mimetype == 'application/pdf'
        bad = client.post(path, data={'expected_reference': old, 'file': (io.BytesIO(b'<svg>unsafe</svg>'), 'bad.png')})
        assert bad.status_code == 400 and current()[field] == old
        replaced = client.post(path, data={'expected_reference': old, 'file': (io.BytesIO(image()), 'replacement.png')})
        assert replaced.status_code == 200 and replaced.json[field] != old
        reference = replaced.json[field]
        preview = client.get(path + '?preview=1')
        assert preview.mimetype == 'image/png' and preview.headers['Content-Disposition'].startswith('inline')
        assert preview.headers['Cache-Control'] == 'private, no-store'
        for method in ['POST', 'DELETE']:
            data = {'expected_reference': old}
            stale = client.post(path, data={**data, 'file': (io.BytesIO(pdf()), 'stale.pdf')}) if method == 'POST' else client.delete(path, json=data)
            assert stale.status_code == 409 and current()[field] == reference
        assert client.delete(path, json={'expected_reference': reference}).status_code == 200
        assert current()[field] is None and client.get(path).status_code == 404
        assert {key: value for key, value in current().items() if key not in ignored} == original
    # The focused row scope works for exports, not just the visible table.
    excel = client.get('/admin/consignments/export.xlsx', query_string={'shipment_id': row['id']})
    assert list(load_workbook(io.BytesIO(excel.data)).active.values)[1][0] == 'FILTER-A1'
    assert load_workbook(io.BytesIO(excel.data)).active.max_row == 2


SCENARIOS = [combined_filters, filtered_exports, filtered_document_zip, library_and_linked_details, invalid_filters, saved_view_identifier_filter, dashboard_missing_documents, individual_document_management]


@pytest.mark.parametrize('scenario', SCENARIOS, ids=lambda fn: fn.__name__)
def test_legacy_filters_documents(admin_client, scenario):
    scenario(admin_client)


@pytest.mark.parametrize('scenario', SCENARIOS, ids=lambda fn: fn.__name__)
def test_er_filters_documents(migrated_admin, scenario):
    scenario(migrated_admin[1])


@pytest.mark.parametrize('path', ['/admin/shipments', '/admin/documents'])
def test_new_pages_require_login(client, path):
    assert client.get(path).status_code == 302
